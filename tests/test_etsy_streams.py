"""The two Etsy income streams, behind Pionir's gate, against fakes at the HTTP boundary.

Pinned here:

- every Etsy listing and Printify publish is PRIVILEGED + ``spends_money`` (the $0.20 listing
  fee) and parks on EVERY call, as its own card, whatever permission the caller holds; a
  Printify product create parks on every call too;
- what Etsy and Printify receive once approved, in their documented request shapes
  (form-urlencoded drafts, multipart files and photos, ``x-api-key: keystring:secret``, the
  OAuth bearer, the token refresh; Printify's JSON with bearer and User-Agent);
- never twice (the ledgers), the daily caps, a file changed after approval is refused, the
  cost-plus rule (reprice, disable, delete), credentials never in a result;
- the receipts reader carries no buyer detail;
- readiness names the exact missing credential, so the workers idle cleanly;
- end to end over REAL Pionir HTTP on a throwaway port: scout -> maker -> parked approval
  -> the Discord card (photo attached, disclosure shown) -> the owner's reaction -> the fake
  Etsy receives the draft, files and photos -> the activation card -> approve -> live; and
  scout -> Printify catalog -> product card -> approve -> publish card -> approve.
"""
from __future__ import annotations

import json
import logging
import tempfile
import threading
import time
import unittest
from http.server import ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from etsy_fakes import (
    ACCESS,
    KEYSTRING,
    NEW_ACCESS,
    NEW_REFRESH,
    PRINTIFY_SHOP,
    PRINTIFY_TOKEN,
    REFRESH,
    SHARED_SECRET,
    SHOP_ID,
    TAXONOMY,
    FakeEtsy,
    FakePrintify,
    search_result,
)
from openpyxl import load_workbook
from test_discord_gate import API, CHANNEL, OWNER
from test_discord_gate import TOKEN as DISCORD_TOKEN
from test_instagram_post import FakeDiscordFiles

from pionir.adapters import etsy as etsy_adapter
from pionir.adapters import printify as printify_adapter
from pionir.adapters.etsy import EtsyAdapter, EtsySettings
from pionir.adapters.printify import PrintifyAdapter, PrintifySettings, net_cents, price_for
from pionir.auth import ensure_tokens, token_path
from pionir.bootstrap import build_runtime
from pionir.config import PionirSettings
from pionir.contracts import RiskLevel, Task
from pionir.crew.etsy import rules, sheets
from pionir.crew.etsy.common import SCOUT_FILE
from pionir.crew.etsy.maker import DigitalMaker
from pionir.crew.etsy.pod import PodMaker
from pionir.crew.etsy.sales import EtsySales
from pionir.crew.etsy.scout import EtsyScout
from pionir.crew.hands import Hands, PionirClient
from pionir.crew.registry import WorkerSpec
from pionir.crew.result import Err, Ok
from pionir.crew.worker import ErrorKind, WorkContext
from pionir.discord_gate import APPROVE, DiscordGate, DiscordGateSettings
from pionir.errors import AdapterProtocolError, AdapterUnavailable
from pionir.server import PionirApp, _make_handler

ETSY = "https://etsy.test/v3"
PRINTIFY = "https://printify.test/v1"
NOW = 1_790_000_000.0

WORDS_DIGITAL = {
    "headline": "Wedding Budget Tracker",
    "title": "Wedding Budget Tracker Spreadsheet | Wedding Planner Printable | Budget "
             "Template",
    "intro": "A simple way for couples to plan what the wedding will cost and to see what is "
             "left as the bills come in.",
    "tags": ["wedding budget", "budget tracker", "wedding planner", "budget spreadsheet",
             "expense tracker"],
    "labels": ["Venue", "Catering", "Flowers and decor", "Photography", "Attire", "Music"],
}
WORDS_POD = {
    "phrase": "Coffee first, then the plan",
    "title": "Coffee First Mug | Funny Coffee Mug | Typography Mug for the Desk",
    "intro": "A simple mug with bold lettering for slow mornings, desk days and anyone who "
             "plans the day around the first cup.",
    "tags": ["coffee mug", "funny mug", "typography mug"],
}


def _settings(root: Path) -> PionirSettings:
    return PionirSettings(
        state_root=root, atani_command=("pionir-test-no-such-binary",),
        daedalus_url=None, melete_url=None, galatea_url=None, crew_url=None,
        bryo_status_command=None, nyx_status_command=None, voodoo_status_command=None,
        embed_model=None, evict_to_fit=False, content_url=None, gumroad_url=None,
        etsy_url=None, printify_url=None)


def spec(worker_id: str, impl: str) -> WorkerSpec:
    division, name = worker_id.split(".")
    return WorkerSpec(worker_id=worker_id, name=name, division=division, impl=impl,
                      kind="etsy", cadence_seconds=3600, provider="none", entities=("Etsy",))


def write_credentials(secrets: Path, *, expires_at: float | None = None) -> Path:
    secrets.mkdir(parents=True, exist_ok=True)
    path = secrets / "etsy.json"
    path.write_text(json.dumps({
        "keystring": KEYSTRING, "shared_secret": SHARED_SECRET, "shop_id": str(SHOP_ID),
        "user_id": "55555", "access_token": ACCESS, "refresh_token": REFRESH,
        "expires_at": expires_at if expires_at is not None else time.time() + 3000}),
        encoding="utf-8")
    (secrets / "printify.json").write_text(json.dumps({
        "token": PRINTIFY_TOKEN, "shop_id": str(PRINTIFY_SHOP)}), encoding="utf-8")
    return path


class _Case(unittest.TestCase):
    """A hermetic Pionir with the Etsy and Printify adapters on the fakes."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self._tmp.name)
        self.secrets = self.root / "secrets"
        self.creds = write_credentials(self.secrets)
        self.etsy_dir = self.root / "etsy"
        self.etsy = FakeEtsy(ETSY)
        self.printify = FakePrintify(PRINTIFY)
        self.etsy_adapter = EtsyAdapter(EtsySettings(
            api_url=ETSY, credentials_file=self.creds, stage_dir=self.etsy_dir / "digital",
            ledger_file=self.etsy_dir / "listings.json", secrets_dir=self.secrets),
            opener=self.etsy)
        self.printify_adapter = PrintifyAdapter(PrintifySettings(
            api_url=PRINTIFY, credentials_file=self.secrets / "printify.json",
            stage_dir=self.etsy_dir / "pod", ledger_file=self.etsy_dir / "printify.json",
            secrets_dir=self.secrets), opener=self.printify)
        runtime = build_runtime(_settings(self.root))
        runtime.register(self.etsy_adapter)
        runtime.register(self.printify_adapter)
        self.app = PionirApp(runtime)

    def tearDown(self) -> None:
        self.app.runtime.cortex.close()
        self._tmp.cleanup()

    # ---- staging a digital listing exactly as the maker does ------------------------------
    def stage(self, slug: str = "budget-wedding-a1b2c3", **over: Any) -> dict:
        maker = DigitalMaker(spec("etsy.digital", "etsy_digital"),
                             stage_dir=str(self.etsy_dir))
        kind = sheets.KINDS["budget"]
        draft, reasons = maker._assemble(dict(WORDS_DIGITAL), kind, "wedding budget sheet")
        self.assertEqual(reasons, [])
        payload = maker._stage(slug, kind, draft, 599, TAXONOMY)
        payload.update(over)
        return payload

    def approve(self, approval_id: str) -> dict:
        res = self.app.approve(approval_id)
        self.assertTrue(res.get("ok"), res)
        self.assertTrue(self.app.jobs.wait(res["task_id"], 60))
        return self.app.approvals.get(approval_id)

    def park(self, capability: str, payload: dict) -> str:
        out = self.app.run_task(capability, payload, permissions=[capability])
        self.assertEqual(out["status"], "pending_approval", out)
        return out["approval_id"]


# ---- the gate ---------------------------------------------------------------------------
class GateTests(_Case):
    def test_every_listing_and_publish_spends_money_and_parks_on_every_call(self) -> None:
        caps = {c.name: c for a in (self.etsy_adapter, self.printify_adapter)
                for c in a.manifest.capabilities}
        for name in ("etsy.create_draft_listing", "etsy.activate_listing", "printify.publish"):
            with self.subTest(capability=name):
                self.assertTrue(caps[name].spends_money)
                self.assertTrue(caps[name].requires_approval)
                self.assertIs(caps[name].risk, RiskLevel.PRIVILEGED)
                self.assertFalse(caps[name].batchable)
                self.assertFalse(caps[name].classifiable)
        self.assertTrue(caps["printify.create_product"].requires_approval)
        self.assertIs(caps["printify.create_product"].risk, RiskLevel.PRIVILEGED)
        for name in ("etsy.search_active", "etsy.receipts", "printify.catalog"):
            self.assertIs(caps[name].risk, RiskLevel.READ_ONLY)
        payload = self.stage()
        for _ in range(3):
            self.park("etsy.create_draft_listing", payload)
        pending = self.app.approvals.pending()
        self.assertEqual(len(pending), 3)
        self.assertFalse(any(p.get("batch") for p in pending))      # never the digest
        self.assertEqual(self.etsy.calls, [])                         # nothing sent

    def test_a_listing_without_the_disclosure_is_refused_before_it_is_parked(self) -> None:
        payload = self.stage()
        payload["description"] = payload["description"].replace(rules.AI_DISCLOSURE, "")
        out = self.app.run_task("etsy.create_draft_listing", payload,
                                permissions=["etsy.create_draft_listing"])
        self.assertNotEqual(out.get("status"), "pending_approval")
        self.assertIn("AI disclosure", json.dumps(out))
        self.assertEqual(self.app.approvals.pending(), [])
        self.assertEqual(self.etsy.calls, [])

    def test_a_denied_listing_never_reaches_etsy(self) -> None:
        aid = self.park("etsy.create_draft_listing", self.stage())
        self.assertTrue(self.app.deny(aid)["ok"])
        self.assertFalse(self.app.approve(aid)["ok"])
        self.assertEqual(self.etsy.calls, [])


class EtsyCallTests(_Case):
    def test_an_approved_draft_sends_etsys_request_shapes(self) -> None:
        payload = self.stage()
        row = self.approve(self.park("etsy.create_draft_listing", payload))
        self.assertEqual(row["status"], "approved", row)
        result = row["result"]["result"]
        self.assertTrue(result["ok"], result)
        self.assertEqual(self.etsy.steps(), ["create", "files", "files", "images", "images",
                                             "images"])
        create = self.etsy.calls[0]
        self.assertEqual((create["method"], create["path"]),
                         ("POST", f"/application/shops/{SHOP_ID}/listings"))
        self.assertEqual(create["api_key"], f"{KEYSTRING}:{SHARED_SECRET}")
        self.assertEqual(create["auth"], f"Bearer {ACCESS}")
        self.assertTrue(create["content_type"].startswith("application/x-www-form-urlencoded"))
        form = create["form"]
        self.assertEqual(form["type"], ["download"])
        self.assertEqual(form["title"], [payload["title"]])
        self.assertEqual(form["description"], [payload["description"]])
        self.assertIn(rules.AI_DISCLOSURE, form["description"][0])
        self.assertEqual(form["price"], ["5.99"])
        self.assertEqual(form["taxonomy_id"], [str(TAXONOMY)])
        self.assertEqual(form["who_made"], ["i_did"])
        self.assertEqual(form["when_made"], ["made_to_order"])
        self.assertEqual(form["should_auto_renew"], ["false"])        # no unapproved renewals
        self.assertEqual(form["tags"], payload["tags"])
        self.assertEqual(len(form["tags"]), 13)
        lid = result["listing_id"]
        stage = self.etsy_dir / "digital" / payload["slug"]
        for call, f in zip(self.etsy.calls[1:3], payload["files"], strict=True):
            self.assertEqual(call["path"], f"/application/shops/{SHOP_ID}/listings/{lid}/files")
            parts = {p["name"]: p for p in call["parts"]}
            self.assertEqual(set(parts), {"file", "name", "rank"})
            self.assertEqual(parts["file"]["filename"], f["name"])
            self.assertEqual(parts["file"]["data"], (stage / f["name"]).read_bytes())
        for rank, (call, im) in enumerate(zip(self.etsy.calls[3:], payload["images"],
                                              strict=True), 1):
            parts = {p["name"]: p for p in call["parts"]}
            self.assertEqual(set(parts), {"image", "rank", "alt_text"})
            self.assertEqual(parts["rank"]["data"], str(rank).encode())
            self.assertEqual(parts["alt_text"]["data"].decode(), im["alt_text"])
            self.assertEqual(parts["image"]["content_type"], "image/png")
        self.assertEqual(self.etsy.listings[lid]["state"], "draft")      # not live yet
        ledger = json.loads((self.etsy_dir / "listings.json").read_text(encoding="utf-8"))
        self.assertEqual(ledger["listings"][payload["slug"]]["state"], "draft")
        # never twice: refused before it is even parked
        out = self.app.run_task("etsy.create_draft_listing", payload,
                                permissions=["etsy.create_draft_listing"])
        self.assertIn("already listed", json.dumps(out))

    def test_activation_only_for_our_own_complete_draft(self) -> None:
        payload = self.stage()
        lid = self.approve(self.park("etsy.create_draft_listing", payload)
                           )["result"]["result"]["listing_id"]
        act = {"slug": payload["slug"], "listing_id": lid, "title": payload["title"],
               "spend": dict(etsy_adapter.LISTING_FEE), "spends_money": True}
        for bad in ({**act, "listing_id": lid + 1}, {**act, "slug": "someone-elses"},
                    {**act, "spends_money": False}, {**act, "spend": {"amount": "0"}}):
            with self.subTest(bad=bad):
                out = self.app.run_task("etsy.activate_listing", bad,
                                        permissions=["etsy.activate_listing"])
                self.assertNotEqual(out.get("status"), "pending_approval", out)
        before = len(self.etsy.calls)
        row = self.approve(self.park("etsy.activate_listing", act))
        self.assertEqual(row["status"], "approved", row)
        read, patch = self.etsy.calls[before:]
        self.assertEqual((read["method"], read["path"]), ("GET", f"/application/listings/{lid}"))
        self.assertEqual((patch["method"], patch["path"]),
                         ("PATCH", f"/application/shops/{SHOP_ID}/listings/{lid}"))
        self.assertEqual(patch["form"], {"state": ["active"]})
        self.assertEqual(self.etsy.listings[lid]["state"], "active")
        self.assertTrue(row["result"]["result"]["url"].startswith("https://www.etsy.com/"))
        out = self.app.run_task("etsy.activate_listing", act,
                                permissions=["etsy.activate_listing"])
        self.assertIn("already active", json.dumps(out))

    def test_a_file_changed_after_parking_is_refused_and_nothing_is_sent(self) -> None:
        payload = self.stage()
        aid = self.park("etsy.create_draft_listing", payload)
        (self.etsy_dir / "digital" / payload["slug"] / payload["files"][0]["name"]) \
            .write_bytes(b"PK\x03\x04 something else")
        row = self.approve(aid)
        self.assertEqual(row["status"], "approved_failed", row)
        self.assertIn("changed since it was checked", json.dumps(row["result"]))
        self.assertEqual(self.etsy.calls, [])

    def test_the_daily_cap_holds(self) -> None:
        settings = EtsySettings(api_url=ETSY, credentials_file=self.creds,
                                stage_dir=self.etsy_dir / "digital",
                                ledger_file=self.etsy_dir / "listings.json",
                                secrets_dir=self.secrets, max_new_listings_per_day=1)
        adapter = EtsyAdapter(settings, opener=self.etsy)
        first, second = self.stage("budget-one-aaaaaa"), self.stage("budget-two-bbbbbb")
        self.assertTrue(adapter.execute(Task("etsy.create_draft_listing", first)).output["ok"])
        with self.assertRaisesRegex(AdapterProtocolError, "daily cap"):
            adapter.validate(Task("etsy.create_draft_listing", second))
        out = adapter.execute(Task("etsy.create_draft_listing", second)).output
        self.assertIn("daily cap", out["error"])
        self.assertEqual(self.etsy.steps().count("create"), 1)

    def test_an_expired_token_is_refreshed_and_saved(self) -> None:
        write_credentials(self.secrets, expires_at=time.time() - 10)
        out = self.etsy_adapter.execute(Task("etsy.receipts", {"min_created": 0})).output
        self.assertTrue(out["ok"], out)
        token, receipts = self.etsy.calls
        self.assertEqual(token["path"], "/public/oauth/token")
        self.assertEqual(token["form"], {"grant_type": ["refresh_token"],
                                         "client_id": [KEYSTRING], "refresh_token": [REFRESH]})
        self.assertEqual(receipts["auth"], f"Bearer {NEW_ACCESS}")
        saved = json.loads(self.creds.read_text(encoding="utf-8"))
        self.assertEqual((saved["access_token"], saved["refresh_token"]),
                         (NEW_ACCESS, NEW_REFRESH))
        self.assertGreater(saved["expires_at"], time.time() + 3000)

    def test_receipts_carry_no_buyer_detail(self) -> None:
        self.etsy.receipts = [{"receipt_id": 1, "name": "A Buyer", "buyer_email": "b@x.test",
                               "first_line": "1 Some Street", "message_from_buyer": "hi",
                               "create_timestamp": NOW, "is_paid": True,
                               "grandtotal": {"amount": 599, "divisor": 100,
                                              "currency_code": "USD"},
                               "transactions": [{"listing_id": 900001, "quantity": 1,
                                                 "title": "x"}]}]
        out = dict(self.etsy_adapter.execute(Task("etsy.receipts", {"min_created": 0})).output)
        self.assertEqual(out["receipts"], [{"receipt_id": 1, "created": NOW, "paid": True,
                                            "grandtotal_cents": 599, "currency": "USD",
                                            "listing_ids": [900001], "items": 1}])
        text = json.dumps(out)
        for private in ("A Buyer", "b@x.test", "Some Street", "hi\""):
            self.assertNotIn(private, text)

    def test_credentials_never_leave_in_a_result_or_a_log(self) -> None:
        self.etsy.fail["create"] = (400, {"error": f"bad key {KEYSTRING}:{SHARED_SECRET} "
                                                   f"token {ACCESS}"})
        records: list = []
        handler = logging.Handler()
        handler.emit = records.append                     # type: ignore[method-assign]
        logging.getLogger().addHandler(handler)
        try:
            out = dict(self.etsy_adapter.execute(Task("etsy.create_draft_listing",
                                                      self.stage())).output)
        finally:
            logging.getLogger().removeHandler(handler)
        self.assertFalse(out["ok"])
        text = json.dumps(out) + "".join(r.getMessage() for r in records)
        for secret in (SHARED_SECRET, ACCESS, REFRESH):
            self.assertNotIn(secret, text)

    def test_missing_credentials_are_named_exactly(self) -> None:
        self.creds.unlink()
        with self.assertRaisesRegex(AdapterUnavailable, r"etsy.json - run tools\\setup-etsy"):
            self.etsy_adapter.validate(Task("etsy.search_active", {"keywords": "budget"}))
        self.creds.write_text(json.dumps({"keystring": KEYSTRING}), encoding="utf-8")
        self.assertIn("has no shared_secret",
                      etsy_adapter.credentials_problem(self.creds))
        scout = EtsyScout(spec("etsy.scout", "etsy_scout"))
        self.assertIn("has no shared_secret", scout.readiness(self.secrets))
        maker = DigitalMaker(spec("etsy.digital", "etsy_digital"), stage_dir=str(self.etsy_dir))
        got = maker.run(WorkContext(now=NOW, http=None, secrets_dir=self.secrets,
                                    state_dir=self.root / "crew", job=lambda j: None,
                                    words=lambda *a: Ok({})))
        self.assertIsInstance(got, Err)
        self.assertEqual(got.error.kind, ErrorKind.NOT_CONFIGURED)
        self.assertIn(r"tools\setup-etsy.ps1", got.error.message)
        (self.secrets / "printify.json").unlink()
        self.assertIn("printify.json", PodMaker(spec("etsy.pod", "etsy_pod")).readiness(
            self.secrets))


class PrintifyTests(_Case):
    def stage_design(self, slug: str = "pod-coffee-mug-abc123", **over: Any) -> dict:
        from pionir.crew.etsy import render
        png = render.design_png(WORDS_POD["phrase"], 2700, 1050)
        folder = self.etsy_dir / "pod" / slug
        folder.mkdir(parents=True)
        (folder / "coffee.png").write_bytes(png)
        import hashlib
        pod = PodMaker(spec("etsy.pod", "etsy_pod"))
        draft, reasons = pod._assemble(dict(WORDS_POD), "funny coffee mug")
        self.assertEqual(reasons, [])
        payload = {"slug": slug, "title": draft["title"], "description": draft["description"],
                   "tags": draft["tags"], "blueprint_id": 5, "print_provider_id": 1,
                   "position": "front",
                   "variants": [{"id": 101, "price_cents": 1799},
                                {"id": 102, "price_cents": 1799}],
                   "design": {"name": "coffee.png", "sha256": hashlib.sha256(png).hexdigest(),
                              "width": 2700, "height": 1050},
                   "min_margin_cents": 400, "max_price_cents": 2399,
                   "ai_disclosure": rules.AI_DISCLOSURE}
        payload.update(over)
        return payload

    def test_cost_plus_reprices_and_disables_and_publish_rechecks(self) -> None:
        self.printify.costs = {101: 700, 102: 1500}
        row = self.approve(self.park("printify.create_product", self.stage_design()))
        result = row["result"]["result"]
        self.assertTrue(result["ok"], result)
        self.assertEqual(self.printify.steps(), ["variants", "upload", "create", "update"])
        create = self.printify.calls[2]
        self.assertEqual(create["path"], f"/shops/{PRINTIFY_SHOP}/products.json")
        self.assertEqual(create["auth"], f"Bearer {PRINTIFY_TOKEN}")
        self.assertTrue(create["agent"])
        body = create["body"]
        self.assertEqual((body["blueprint_id"], body["print_provider_id"]), (5, 1))
        self.assertEqual(body["print_areas"][0]["placeholders"][0]["images"][0]["id"],
                         "img0001")
        self.assertIn(rules.AI_DISCLOSURE, body["description"])
        # 101 keeps the margin at 17.99; 102 (cost 15.00) is repriced to keep 4.00
        self.assertEqual(result["repriced"], [102])
        prices = {v["id"]: v["price_cents"] for v in result["variants"]}
        self.assertEqual(prices[101], 1799)
        self.assertEqual(prices[102], price_for(1500, 400))
        self.assertGreaterEqual(net_cents(prices[102], 1500), 400)
        self.assertTrue(str(prices[102]).endswith("99"))
        pid = result["product_id"]
        # publish: the card's prices must be the live ones
        pub = {"slug": "pod-coffee-mug-abc123", "product_id": pid, "title": "x",
               "variants": [{k: v[k] for k in ("id", "price_cents", "cost_cents")}
                            for v in result["variants"]],
               "min_margin_cents": 400, "spend": dict(printify_adapter.LISTING_FEE),
               "spends_money": True}
        self.printify.products[pid]["variants"][0]["cost"] = 900     # the cost moved
        row = self.approve(self.park("printify.publish", pub))
        self.assertEqual(row["status"], "approved_failed")
        self.assertEqual(self.printify.published, [])
        self.printify.products[pid]["variants"][0]["cost"] = 700
        row = self.approve(self.park("printify.publish", pub))
        self.assertEqual(row["status"], "approved", row)
        self.assertEqual(self.printify.published, [pid])
        publish = self.printify.calls[-1]
        self.assertEqual(publish["body"], {"title": True, "description": True, "images": True,
                                           "variants": True, "tags": True,
                                           "keyFeatures": True, "shipping_template": True})

    def test_a_product_that_cannot_keep_the_margin_is_deleted(self) -> None:
        self.printify.costs = {101: 3000, 102: 3000}
        row = self.approve(self.park("printify.create_product", self.stage_design()))
        self.assertEqual(row["status"], "approved_failed")
        self.assertIn("deleted", json.dumps(row["result"]))
        self.assertEqual(self.printify.products, {})
        self.assertIn("delete", self.printify.steps())

    def test_a_design_not_the_print_area_size_is_refused(self) -> None:
        payload = self.stage_design()
        payload["design"] = {**payload["design"], "width": 2000}
        out = self.app.run_task("printify.create_product", payload,
                                permissions=["printify.create_product"])
        self.assertNotEqual(out.get("status"), "pending_approval")
        self.assertEqual(self.printify.calls, [])


# ---- end to end, over real Pionir HTTP -------------------------------------------------------
class EndToEndTests(_Case):
    def setUp(self) -> None:
        super().setUp()
        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), _make_handler(self.app))
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()
        tokens = self.app.runtime.settings.client_token_path
        ensure_tokens(tokens, ("crew",))
        client = PionirClient(f"http://127.0.0.1:{self.httpd.server_address[1]}",
                              timeout=30, token_file=token_path(tokens, "crew"))
        self.hands = Hands(SimpleNamespace(job_follow_seconds=30, job_poll_seconds=0.2),
                           None, client)
        self.state = self.root / "crew" / "workers"
        self.state.mkdir(parents=True)
        self.etsy.default_search = self._search

    def tearDown(self) -> None:
        self.httpd.shutdown()
        self.httpd.server_close()
        super().tearDown()

    @staticmethod
    def _search(keyword: str) -> dict:
        # the budget seed has the most favourites per competing listing; its results'
        # tags carry "wedding budget sheet", which the next scout run measures
        fav = {"budget tracker spreadsheet": 90, "funny coffee mug": 70}.get(keyword, 20)
        return search_result(keyword, base_fav=fav,
                             tags=("wedding budget sheet", "budget planner", "coffee lover"))

    def ctx(self, words: dict | None = None) -> WorkContext:
        return WorkContext(now=time.time(), http=None, secrets_dir=self.secrets,
                           state_dir=self.state, job=self.hands._run_job,
                           approval=self.hands.approval,
                           words=(lambda *a: Ok(dict(words))) if words else None)

    def gate(self, fake: FakeDiscordFiles) -> DiscordGate:
        token_file = self.root / "discord-bot-token.txt"
        token_file.write_text(DISCORD_TOKEN, encoding="utf-8")
        return DiscordGate.for_app(
            self.app, DiscordGateSettings(state_root=self.root / "gate", channel_id=CHANNEL,
                                          owner_user_id=OWNER, token_file=token_file,
                                          api_base=API, poll_seconds=0.01),
            opener=fake, sleep=lambda _s: None)

    def scout(self) -> dict:
        scout = EtsyScout(spec("etsy.scout", "etsy_scout"), max_probes=8)
        got = scout.run(self.ctx())
        self.assertIsInstance(got, Ok, got)
        return json.loads((self.state / SCOUT_FILE).read_text(encoding="utf-8"))

    def test_digital_scout_to_live_listing_through_the_owners_reactions(self) -> None:
        ranking = self.scout()["ranking"]
        searched = [c["query"]["keywords"][0] for c in self.etsy.calls
                    if c.get("step") == "search"]
        self.assertIn("budget tracker spreadsheet", searched)        # seeds: nothing measured
        best = next(r for r in ranking if r["track"] == "digital")
        self.assertEqual(best["keyword"], "budget tracker spreadsheet")
        self.assertEqual(best["taxonomy_id"], TAXONOMY)
        # the second run measures what Etsy's own tags suggested, not the seeds again
        self.scout()
        searched2 = [c["query"]["keywords"][0] for c in self.etsy.calls
                     if c.get("step") == "search"][len(searched):]
        self.assertIn("wedding budget sheet", searched2)
        self.assertNotIn("budget tracker spreadsheet", searched2)

        maker = DigitalMaker(spec("etsy.digital", "etsy_digital"),
                             stage_dir=str(self.etsy_dir), daily_cap=1)
        got = maker.run(self.ctx(WORDS_DIGITAL))
        self.assertIsInstance(got, Ok, got)
        pending = self.app.approvals.pending()
        self.assertEqual([p["capability"] for p in pending], ["etsy.create_draft_listing"])
        self.assertEqual(self.etsy.steps().count("create"), 0)        # parked, not sent
        payload = pending[0]["payload"]
        # the staged workbook is a real one, with formulas, the one the card pins
        folder = self.etsy_dir / "digital" / payload["slug"]
        xlsx = folder / next(f["name"] for f in payload["files"] if f["name"].endswith(".xlsx"))
        wb = load_workbook(xlsx)
        self.assertEqual(wb["Tracker"]["D6"].value, "=B6-C6")
        self.assertEqual(wb["Tracker"]["A6"].value, "Venue")

        # the Discord card: title, price, disclosure, files, photos, the photo attached
        fake = FakeDiscordFiles()
        gate = self.gate(fake)
        self.assertTrue(gate.run_once())
        (payload_part, file_part), = fake.uploads
        self.assertEqual(file_part["content_type"], "image/png")
        self.assertEqual(file_part["data"], (folder / "1-example.png").read_bytes())
        text = "\n".join(p["content"] for p in fake.posts())
        self.assertIn(etsy_adapter.ETSY_LINE, text)
        self.assertIn("SPENDS MONEY: 0.20 USD", text)
        self.assertIn(f"**Title:** {payload['title']}", text)
        self.assertIn("**Price:** 6.99 USD", text)      # the measured median, x.99
        self.assertIn("**AI disclosure in the description:** yes, word for word", text)
        self.assertIn(rules.AI_DISCLOSURE, text)
        self.assertIn(".xlsx`", text)
        self.assertIn("-printable.pdf`", text)
        self.assertIn("`1-example.png` 2000x1500", text)

        # the owner's ✅ in Discord runs it: the fake Etsy receives the draft and its files
        state = json.loads(gate.settings.state_path.read_text(encoding="utf-8"))
        fake.react(state["messages"][pending[0]["id"]]["message_id"], APPROVE, OWNER)
        gate.run_once()
        row = self.app.approvals.get(pending[0]["id"])
        self.assertTrue(self.app.jobs.wait(row["task_id"], 60))
        row = self.app.approvals.get(pending[0]["id"])
        self.assertEqual(row["status"], "approved", row)
        self.assertEqual(self.etsy.steps()[-6:], ["create", "files", "files", "images",
                                                  "images", "images"])
        lid = row["result"]["result"]["listing_id"]
        self.assertEqual(self.etsy.listings[lid]["state"], "draft")

        # the maker sees the draft and asks for the second yes: going live
        got = maker.run(self.ctx(WORDS_DIGITAL))
        self.assertIsInstance(got, Ok, got)
        act = [p for p in self.app.approvals.pending()
               if p["capability"] == "etsy.activate_listing"]
        self.assertEqual(len(act), 1)
        self.assertEqual(act[0]["payload"]["listing_id"], lid)
        self.approve(act[0]["id"])
        self.assertEqual(self.etsy.listings[lid]["state"], "active")
        got = maker.run(self.ctx(WORDS_DIGITAL))
        kinds = [o.kind for o in got.value]
        self.assertIn("etsy.listed", kinds)
        rec = json.loads((self.state / "etsy.digital.json").read_text(encoding="utf-8"))
        # the daily cap (1 here) held: nothing new was made the same day
        self.assertEqual([i["status"] for i in rec["items"]], ["active"])
        self.assertEqual(got.value[-1].payload["note"], "the daily cap of 1 is reached")
        # dedupe: the keyword is recorded as made, for ever
        self.assertEqual(list(rec["made"]), ["budget tracker spreadsheet"])
        # the next day, the next keyword - never the same one again
        tomorrow = self.ctx(WORDS_DIGITAL)
        tomorrow.now += 86400
        maker.run(tomorrow)
        rec = json.loads((self.state / "etsy.digital.json").read_text(encoding="utf-8"))
        self.assertEqual(len(rec["made"]), 2)
        self.assertNotEqual(rec["items"][1]["keyword"], "budget tracker spreadsheet")

        # sales: read-only, amounts only
        self.etsy.receipts = [{"receipt_id": 7, "create_timestamp": time.time(),
                               "is_paid": True, "name": "Private Person",
                               "grandtotal": {"amount": 599, "divisor": 100,
                                              "currency_code": "USD"},
                               "transactions": [{"listing_id": lid, "quantity": 1}]}]
        sales = EtsySales(spec("treasury.etsy", "etsy_sales")).run(self.ctx())
        self.assertIsInstance(sales, Ok, sales)
        tally = sales.value[-1]
        gross = next(f for f in tally.figures if f.window == "last30" and
                     f.unit == "usd_cents")
        self.assertEqual(gross.value, 599)
        self.assertNotIn("Private Person", json.dumps([o.payload for o in sales.value]))

    def test_print_on_demand_scout_to_published_through_two_approvals(self) -> None:
        ranking = self.scout()["ranking"]
        best = next(r for r in ranking if r["track"] == "pod")
        self.assertEqual((best["keyword"], best["product"]), ("funny coffee mug", "mug"))
        pod = PodMaker(spec("etsy.pod", "etsy_pod"), stage_dir=str(self.etsy_dir))
        got = pod.run(self.ctx(WORDS_POD))
        self.assertIsInstance(got, Ok, got)
        # the catalog read ran directly (read-only); the create is parked
        self.assertEqual(self.printify.steps()[:5], ["blueprints", "providers", "variants",
                                                     "providers", "variants"])
        (create,) = self.app.approvals.pending()
        self.assertEqual(create["capability"], "printify.create_product")
        p = create["payload"]
        self.assertEqual((p["blueprint_id"], p["print_provider_id"]), (5, 1))
        self.assertEqual([v["id"] for v in p["variants"]], [101, 102])     # light only
        self.assertEqual((p["design"]["width"], p["design"]["height"]), (2700, 1050))
        fake = FakeDiscordFiles()
        gate = self.gate(fake)
        gate.run_once()
        self.assertEqual(len(fake.uploads), 1)                          # the design attached
        self.approve(create["id"])
        self.assertIn("create", self.printify.steps())
        got = pod.run(self.ctx(WORDS_POD))
        self.assertIsInstance(got, Ok, got)
        (publish,) = [r for r in self.app.approvals.pending()
                      if r["capability"] == "printify.publish"]
        text_rows = [v for v in publish["payload"]["variants"]]
        self.assertTrue(all(net_cents(v["price_cents"], v["cost_cents"]) >= 400
                            for v in text_rows))
        self.assertEqual(self.printify.published, [])
        self.approve(publish["id"])
        self.assertEqual(len(self.printify.published), 1)
        got = pod.run(self.ctx(WORDS_POD))
        self.assertIn("etsy.pod_published", [o.kind for o in got.value])


if __name__ == "__main__":
    unittest.main()
