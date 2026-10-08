"""The Etsy streams' products: the fail-closed listing rules, the workbooks and the pictures.

- A listing without the AI disclosure word for word is refused (by the rules, by the
  adapter's check, and by Pionir before anything is parked).
- Brands, proper nouns, unknown words and claims are refused; exactly 13 Etsy-shaped tags.
- Every workbook kind opens with openpyxl, carries its formulas in the cells they belong
  in, referencing the cells they should, and those formulas compute the right numbers.
- The printable PDF and the listing photos are drawn from the workbook itself.
"""
from __future__ import annotations

import io
import unittest
import zipfile

from openpyxl import load_workbook

from pionir.adapters import etsy as etsy_adapter
from pionir.crew.etsy import render, rules, sheets
from pionir.crew.etsy.formulas import Evaluator, FormulaError, references
from pionir.crew.etsy.maker import describe, normalise_tags

TAGS = ["budget tracker", "budget spreadsheet", "wedding budget", "wedding planner",
        "budget planner", "expense tracker", "printable planner", "spending tracker",
        "money planner", "monthly budget", "digital download", "planner printable",
        "budget template"]
TITLE = "Wedding Budget Tracker Spreadsheet | Wedding Planner Printable | Budget Template"
INTRO = ("A simple way for couples to plan what the wedding will cost and to see what is "
         "left as the bills come in.")
LABELS = ["Venue", "Catering", "Flowers and decor", "Photography", "Attire", "Music"]


def description(intro: str = INTRO) -> str:
    return describe(intro, sheets.KINDS["budget"], LABELS)


class DisclosureTests(unittest.TestCase):
    def test_a_listing_without_the_ai_disclosure_is_refused(self) -> None:
        good = description()
        self.assertEqual(rules.listing_problems(TITLE, good, TAGS, kind="digital"), [])
        bare = good.replace(rules.AI_DISCLOSURE, "")
        problems = rules.listing_problems(TITLE, bare, TAGS, kind="digital")
        self.assertTrue(any("AI disclosure" in p for p in problems), problems)
        # a disclosure reworded is not the disclosure
        reworded = good.replace(rules.AI_DISCLOSURE, rules.AI_DISCLOSURE.replace("AI", "A.I."))
        self.assertTrue(rules.listing_problems(TITLE, reworded, TAGS, kind="digital"))
        with self.assertRaisesRegex(ValueError, "AI disclosure"):
            rules.check_listing(TITLE, bare, TAGS, kind="digital")

    def test_the_digital_and_pod_notes_are_required_too(self) -> None:
        good = description()
        no_note = good.replace(rules.DIGITAL_NOTE, "")
        self.assertTrue(any("required note" in p for p in
                            rules.listing_problems(TITLE, no_note, TAGS, kind="digital")))
        # a digital description is not a print-on-demand one
        self.assertTrue(rules.listing_problems(TITLE, good, TAGS, kind="pod"))

    def test_the_adapter_check_refuses_it_too(self) -> None:
        payload = {"slug": "budget-x-1a2b3c", "title": TITLE,
                   "description": description().replace(rules.AI_DISCLOSURE, ""),
                   "tags": TAGS, "price_cents": 599, "taxonomy_id": 6844,
                   "files": [{"name": "a.xlsx", "sha256": "0" * 64}],
                   "images": [{"name": "1.png", "sha256": "0" * 64, "alt_text": "a photo"}],
                   "ai_disclosure": rules.AI_DISCLOSURE,
                   "spend": dict(etsy_adapter.LISTING_FEE), "spends_money": True}
        with self.assertRaisesRegex(ValueError, "AI disclosure"):
            etsy_adapter.check_create(payload)
        with self.assertRaisesRegex(ValueError, "ai_disclosure"):
            etsy_adapter.check_create({**payload, "description": description(),
                                       "ai_disclosure": "made with AI"})


class TextRuleTests(unittest.TestCase):
    def test_brands_are_refused_in_any_case(self) -> None:
        for text in ("works in Google Sheets", "an excel budget", "for your Cricut",
                     "goodnotes planner", "a Disney trip planner", "Target run list",
                     "Taylor Swift lyrics", "notion template pack"):
            with self.subTest(text=text):
                self.assertIsNotNone(rules.brand_problem(text))
        self.assertIsNone(rules.brand_problem("an apple pie recipe card"))   # the fruit
        self.assertIsNone(rules.brand_problem("a target amount for each goal"))

    def test_proper_nouns_and_unknown_words_are_refused(self) -> None:
        for text in ("A budget for Kimberly and her family", "Plan a trip to Seattle",
                     "made for Verizon customers", "the zorbleflax tracker",
                     "NASA style planner"):
            with self.subTest(text=text):
                self.assertIsNotNone(rules.name_problem(text))
        for text in ("Mark each habit as done. Plan your Monday.", "A PDF and an XLSX file",
                     "Wedding Budget Tracker", "the totals work themselves out"):
            with self.subTest(text=text):
                self.assertIsNone(rules.name_problem(text))

    def test_claims_links_and_contacts_are_refused(self) -> None:
        for text in ("guaranteed to save you money", "the best seller on the shop",
                     "lose weight with this tracker", "see www.example.com",
                     "write to me at a@b.co", "call 555-123-4567", "100% satisfaction",
                     "a #budget tracker", "made by @someone"):
            with self.subTest(text=text):
                self.assertTrue(rules.text_problems("t", text, multiline=True), text)

    def test_tags_are_exactly_thirteen_in_etsys_form(self) -> None:
        self.assertEqual(rules.tag_problems(TAGS), [])
        self.assertTrue(rules.tag_problems(TAGS[:12]))
        self.assertTrue(rules.tag_problems([*TAGS[:12], "this tag is far too long"]))
        self.assertTrue(rules.tag_problems([*TAGS[:12], "Budget"]))
        self.assertTrue(rules.tag_problems([*TAGS[:12], TAGS[0]]))
        self.assertTrue(rules.tag_problems([*TAGS[:12], "canva template"]))

    def test_title_rules(self) -> None:
        self.assertEqual(rules.title_problems(TITLE), [])
        self.assertTrue(rules.title_problems("x" * 141))
        self.assertTrue(rules.title_problems("Budget: plan: spend tracker sheet"))
        self.assertTrue(rules.title_problems("Budget tracker for $5 a month"))

    def test_normalised_tags_are_topped_up_to_thirteen_and_still_checked(self) -> None:
        kind = sheets.KINDS["budget"]
        tags = normalise_tags(["Wedding Budget!!", "an excel sheet", "x" * 30], kind,
                              "budget tracker spreadsheet")
        self.assertEqual(len(tags), 13)
        self.assertEqual(rules.tag_problems(tags), [])
        self.assertIn("wedding budget", tags)
        self.assertNotIn("an excel sheet", tags)          # a brand never gets in

    def test_keywords_are_never_brands(self) -> None:
        self.assertEqual(rules.keyword_problems("budget tracker"), [])
        self.assertTrue(rules.keyword_problems("disney planner"))
        self.assertTrue(rules.keyword_problems("Budget Tracker"))

    def test_the_fixed_text_passes_its_own_rules(self) -> None:
        for kind in sheets.KINDS.values():
            for line in kind.howto:
                with self.subTest(kind=kind.name, line=line):
                    self.assertEqual(rules.text_problems("howto", line), [])
        for sentence in (rules.AI_DISCLOSURE, rules.DIGITAL_NOTE, rules.POD_NOTE):
            self.assertEqual(rules.text_problems("note", sentence), [])


# ---- the workbooks -------------------------------------------------------------------------
LABELS4 = ["Rent", "Food", "Travel", "Phone"]


def _wb(kind: str, labels=LABELS4):
    data = sheets.build_workbook(kind, "Monthly Test Tracker", labels)
    return data, load_workbook(io.BytesIO(data))


def _values(data: bytes, sheet: str) -> dict:
    """Every cell's value with formulas computed by our evaluator."""
    return sheets.read_sheet(data, sheet).values


class WorkbookTests(unittest.TestCase):
    def test_every_kind_opens_with_three_sheets_and_no_macros(self) -> None:
        for kind in sheets.KINDS:
            with self.subTest(kind=kind):
                data, wb = _wb(kind)
                self.assertEqual(wb.sheetnames, ["Tracker", "Example", "How to use"])
                names = zipfile.ZipFile(io.BytesIO(data)).namelist()
                self.assertFalse([n for n in names if "vba" in n.lower()
                                  or "externallink" in n.lower()])

    def test_formulas_are_in_place_in_both_sheets_and_read_the_right_cells(self) -> None:
        for kind_name, kind in sheets.KINDS.items():
            data, wb = _wb(kind_name)
            expected = sheets.formulas_of(kind, len(LABELS4))
            self.assertTrue(expected)
            for sheet in ("Tracker", "Example"):
                ws = wb[sheet]
                for cell, formula in expected.items():
                    with self.subTest(kind=kind_name, sheet=sheet, cell=cell):
                        self.assertEqual(ws[cell].value, formula)
                        row = int("".join(c for c in cell if c.isdigit()))
                        lay = sheets.layout_of(kind, len(LABELS4))
                        refs = references(formula)
                        self.assertTrue(refs)
                        for ref in refs:
                            r = int("".join(c for c in ref if c.isdigit()))
                            # a row formula reads its own row (or a top input above the
                            # table); a total reads the data rows or its own total row
                            if lay.first <= row <= lay.last:
                                self.assertTrue(r == row or r < lay.header_row, (cell, ref))
                            elif row == lay.total:
                                self.assertTrue(lay.first <= r <= lay.total
                                                or r < lay.header_row, (cell, ref))
                        # the computed sheets evaluate every formula without an error
                        Evaluator(lambda ref, ws=ws: ws[ref].value).value(cell)

    def test_the_budget_computes(self) -> None:
        data, wb = _wb("budget")
        ws = wb["Example"]
        v = _values(data, "Example")
        lay = sheets.layout_of(sheets.KINDS["budget"], 4)
        planned = [ws[f"B{r}"].value for r in range(lay.first, lay.last + 1)]
        actual = [ws[f"C{r}"].value for r in range(lay.first, lay.last + 1)]
        for i, r in enumerate(range(lay.first, lay.last + 1)):
            self.assertEqual(v[f"D{r}"], planned[i] - actual[i])
            self.assertAlmostEqual(v[f"E{r}"], actual[i] / planned[i])
        self.assertEqual(v[f"B{lay.total}"], sum(planned))
        self.assertEqual(v[f"C{lay.total}"], sum(actual))
        self.assertEqual(v[f"D{lay.total}"], sum(planned) - sum(actual))
        self.assertEqual(v["B3"], ws["B2"].value - sum(actual))
        # the buyer's blank sheet computes to nothing, not to an error
        t = _values(data, "Tracker")
        self.assertEqual(t[f"E{lay.first}"], "")
        self.assertEqual(t[f"D{lay.total}"], 0)

    def test_the_habit_tracker_counts_marks(self) -> None:
        data, wb = _wb("habit")
        ws = wb["Example"]
        v = _values(data, "Example")
        lay = sheets.layout_of(sheets.KINDS["habit"], 4)
        days = ws["B2"].value
        total = 0
        for r in range(lay.first, lay.last + 1):
            marks = sum(1 for c in range(2, 33) if ws.cell(r, c).value == "x")
            total += marks
            self.assertEqual(v[f"AG{r}"], marks)
            self.assertAlmostEqual(v[f"AH{r}"], marks / days)
        self.assertEqual(v[f"AG{lay.total}"], total)
        self.assertAlmostEqual(v[f"AH{lay.total}"], total / (days * 4))

    def test_savings_and_debt_compute(self) -> None:
        data, wb = _wb("savings")
        ws, v = wb["Example"], _values(data, "Example")
        lay = sheets.layout_of(sheets.KINDS["savings"], 4)
        for r in range(lay.first, lay.last + 1):
            target, saved = ws[f"B{r}"].value, ws[f"C{r}"].value
            self.assertEqual(v[f"D{r}"], max(target - saved, 0))
            self.assertAlmostEqual(v[f"E{r}"], min(saved / target, 1))
        data, wb = _wb("debt")
        ws, v = wb["Example"], _values(data, "Example")
        lay = sheets.layout_of(sheets.KINDS["debt"], 4)
        for r in range(lay.first, lay.last + 1):
            bal, rate, pay = (ws[f"B{r}"].value, ws[f"C{r}"].value, ws[f"D{r}"].value)
            interest = round(bal * rate / 12, 2)
            self.assertAlmostEqual(v[f"E{r}"], interest)
            self.assertAlmostEqual(v[f"F{r}"], max(bal + interest - pay, 0))

    def test_the_evaluator_refuses_what_it_does_not_know(self) -> None:
        ev = Evaluator(lambda ref: None)
        with self.assertRaises(FormulaError):
            ev.evaluate("=VLOOKUP(A1,B1:C3,2)")
        with self.assertRaises(FormulaError):
            Evaluator(lambda ref: "=A1").value("A1")      # circular

    def test_keywords_map_to_kinds_only_when_we_can_make_them(self) -> None:
        self.assertEqual(sheets.kind_for("wedding budget sheet"), "budget")
        self.assertEqual(sheets.kind_for("habit tracker printable"), "habit")
        self.assertEqual(sheets.kind_for("debt payoff tracker"), "debt")
        self.assertIsNone(sheets.kind_for("wedding budget"))          # names no format
        self.assertIsNone(sheets.kind_for("funny coffee mug"))


class PictureTests(unittest.TestCase):
    def test_the_printable_and_photos_are_drawn_from_the_workbook(self) -> None:
        from PIL import Image

        data = sheets.build_workbook("budget", "Wedding Budget Tracker", LABELS)
        pages = render.printable_pages(data, "budget")
        pdf = render.pdf_bytes(pages)
        self.assertTrue(pdf.startswith(b"%PDF-"))
        photos = render.preview_images(data, pages, "budget", "Wedding Budget Tracker",
                                       [("a.xlsx", len(data)), ("a.pdf", len(pdf))])
        self.assertEqual([n for n, _b, _a in photos],
                         ["1-example.png", "2-printable.png", "3-included.png"])
        for name, png, alt in photos:
            img = Image.open(io.BytesIO(png))
            self.assertEqual(img.size, render.PREVIEW_SIZE, name)
            self.assertGreaterEqual(img.size[0], etsy_adapter.MIN_IMAGE_WIDTH)
            self.assertEqual(rules.text_problems("alt", alt), [])
        # the printable page carries the workbook's own labels: change a label and the
        # page changes
        other = sheets.build_workbook("budget", "Wedding Budget Tracker",
                                      ["Cake", *LABELS[1:]])
        self.assertNotEqual(render.printable_pages(other, "budget")[0].tobytes(),
                            pages[0].tobytes())

    def test_a_design_is_the_exact_print_area_and_transparent(self) -> None:
        from PIL import Image

        png = render.design_png("Coffee first, then the plan", 2700, 1050)
        img = Image.open(io.BytesIO(png))
        self.assertEqual(img.size, (2700, 1050))
        self.assertEqual(img.mode, "RGBA")
        self.assertEqual(img.getpixel((0, 0))[3], 0)


if __name__ == "__main__":
    unittest.main()
