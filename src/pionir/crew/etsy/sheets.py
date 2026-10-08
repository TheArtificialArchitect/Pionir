"""The spreadsheets the digital maker can build: real .xlsx workbooks with working formulas.

A ``Kind`` is something we can genuinely make, not a niche: the scout measures demand for
keywords, and a keyword is made only when its words point at a kind (``kind_for``). Each
workbook has three sheets:

- **Tracker** - the buyer's own: row labels (written by the model for the niche, checked),
  blank inputs, and formulas that compute as soon as numbers go in;
- **Example** - the same layout and formulas with sample numbers, so the buyer sees it
  work (and the listing's preview shows exactly this sheet, computed);
- **How to use** - plain instructions, fixed text written here (checked like a listing).

Formulas are the ``formulas`` module's subset (SUM, COUNTIF, IF, IFERROR, ROUND, MAX, MIN
and arithmetic), so every spreadsheet app computes them and the tests can prove they do.
``openpyxl`` is imported lazily: a crew without it reports the maker NOT CONFIGURED rather
than failing to load.
"""
from __future__ import annotations

import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field

from .formulas import Evaluator, col_letters

TITLE_ROW = 1
MONEY = "#,##0.00"
PERCENT = "0%"
WHOLE = "0"
SHEETS = ("Tracker", "Example", "How to use")
# Words in a keyword that say "a file you fill in", for a digital listing.
FORMAT_WORDS = frozenset({"spreadsheet", "spreadsheets", "tracker", "trackers", "template",
                          "templates", "planner", "planners", "printable", "printables",
                          "sheet", "sheets", "worksheet", "worksheets", "log", "calculator",
                          "chart", "checklist", "journal", "digital", "download"})


class OpenpyxlMissing(RuntimeError):
    pass


def openpyxl_problem() -> str | None:
    """Why the maker cannot write workbooks, or None."""
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return ("openpyxl is not installed in the crew's Python - run: "
                "python -m pip install openpyxl")
    return None


@dataclass(frozen=True)
class Column:
    header: str
    role: str                         # "label" | "input" | "formula"
    formula: str = ""                 # with {r} (this row); a formula column
    total: str = ""                   # the total row's formula, {f} {l} {t} {n}
    number_format: str = ""
    width: float = 14
    sample: Callable[[int], object] | None = None   # Example sheet value for row index i


@dataclass(frozen=True)
class TopCell:
    label: str
    cell: str                          # e.g. "B2"; its label goes in column A of that row
    role: str                          # "input" | "formula"
    value: object = None               # a default input (the Tracker gets it too)
    sample: object = None              # the Example sheet's value
    formula: str = ""                  # with {t} (the total row)
    number_format: str = ""


@dataclass(frozen=True)
class Kind:
    name: str
    noun: str                          # "budget tracker"
    label_word: str                    # what the model names: "spending category"
    vocabulary: frozenset
    columns: tuple
    top: tuple = ()
    labels: tuple = (4, 12)            # how many row labels
    howto: tuple = ()                  # the How to use sheet, one line each
    landscape: bool = False
    pod: bool = False                  # never; kept for symmetry with the scout's split
    extra: dict = field(default_factory=dict)


def _days() -> tuple:
    return tuple(Column(str(d), "input", width=3.2,
                        sample=(lambda i, d=d: "x" if (d * 7 + i * 3) % 5 < 3 else None))
                 for d in range(1, 32))


KINDS: dict[str, Kind] = {
    "budget": Kind(
        "budget", "budget tracker", "spending category",
        frozenset({"budget", "budgeting", "expense", "expenses", "spending", "bills", "bill",
                   "finance", "finances", "money", "household", "paycheck"}),
        columns=(
            Column("Category", "label", width=26),
            Column("Planned", "input", total="=SUM(B{f}:B{l})", number_format=MONEY,
                   sample=lambda i: 120 + 45 * ((i * 7) % 9)),
            Column("Actual", "input", total="=SUM(C{f}:C{l})", number_format=MONEY,
                   sample=lambda i: 110 + 45 * ((i * 5) % 9)),
            Column("Difference", "formula", formula="=B{r}-C{r}", total="=B{t}-C{t}",
                   number_format=MONEY),
            Column("Share of plan used", "formula", formula='=IF(B{r}=0,"",C{r}/B{r})',
                   total='=IF(B{t}=0,"",C{t}/B{t})', number_format=PERCENT, width=18),
        ),
        top=(TopCell("Monthly income", "B2", "input", sample=3200, number_format=MONEY),
             TopCell("Left after actual spending", "B3", "formula", formula="=B2-C{t}",
                     number_format=MONEY)),
        howto=("Type your monthly income in the cell next to Monthly income.",
               "Fill in the Planned column for each category before the month starts.",
               "Fill in the Actual column as you spend.",
               "Difference, Share of plan used and the totals work themselves out.",
               "The Example sheet shows the same layout with sample numbers.",
               "Rename any category by typing over it."),
    ),
    "habit": Kind(
        "habit", "habit tracker", "habit to track",
        frozenset({"habit", "habits", "routine", "routines", "daily", "streak", "goals"}),
        columns=(Column("Habit", "label", width=24), *_days(),
                 Column("Days done", "formula", formula='=COUNTIF(B{r}:AF{r},"x")',
                        total="=SUM(AG{f}:AG{l})", number_format=WHOLE, width=11),
                 Column("Rate", "formula", formula='=IF($B$2=0,"",AG{r}/$B$2)',
                        total='=IF($B$2=0,"",AG{t}/($B$2*{n}))', number_format=PERCENT,
                        width=8)),
        top=(TopCell("Days in this month", "B2", "input", value=31, sample=30,
                     number_format=WHOLE),),
        labels=(4, 10),
        howto=("Set the number of days in the month next to Days in this month.",
               "Type a lower-case x under each day you kept a habit.",
               "Days done counts the x marks; Rate divides by the days in the month.",
               "The Example sheet shows a filled-in month.",
               "Rename any habit by typing over it."),
        landscape=True,
    ),
    "savings": Kind(
        "savings", "savings goal tracker", "savings goal",
        frozenset({"savings", "saving", "save", "sinking", "fund", "funds", "emergency",
                   "goal", "challenge"}),
        columns=(
            Column("Goal", "label", width=26),
            Column("Goal amount", "input", total="=SUM(B{f}:B{l})", number_format=MONEY,
                   sample=lambda i: 500 * (1 + (i * 3) % 7)),
            Column("Saved so far", "input", total="=SUM(C{f}:C{l})", number_format=MONEY,
                   sample=lambda i: 150 * (1 + (i * 5) % 7)),
            Column("Still to save", "formula", formula="=MAX(B{r}-C{r},0)",
                   total="=SUM(D{f}:D{l})", number_format=MONEY),
            Column("Progress", "formula", formula='=IF(B{r}=0,"",MIN(C{r}/B{r},1))',
                   total='=IF(B{t}=0,"",MIN(C{t}/B{t},1))', number_format=PERCENT),
        ),
        howto=("Type the amount you want to reach for each goal in the Goal amount column.",
               "Update Saved so far whenever you put money aside.",
               "Still to save and Progress work themselves out.",
               "The Example sheet shows the same layout with sample numbers.",
               "Rename any goal by typing over it."),
    ),
    "debt": Kind(
        "debt", "debt payoff tracker", "debt to pay off",
        frozenset({"debt", "debts", "payoff", "loan", "loans", "snowball", "avalanche",
                   "credit"}),
        columns=(
            Column("Debt", "label", width=24),
            Column("Balance", "input", total="=SUM(B{f}:B{l})", number_format=MONEY,
                   sample=lambda i: 800 + 650 * ((i * 4) % 6)),
            Column("Yearly interest rate", "input", number_format=PERCENT, width=18,
                   sample=lambda i: round(0.04 + 0.03 * ((i * 3) % 5), 2)),
            Column("Monthly payment", "input", total="=SUM(D{f}:D{l})", number_format=MONEY,
                   width=16, sample=lambda i: 60 + 25 * ((i * 2) % 5)),
            Column("Interest this month", "formula", formula="=ROUND(B{r}*C{r}/12,2)",
                   total="=SUM(E{f}:E{l})", number_format=MONEY, width=18),
            Column("Balance after payment", "formula", formula="=MAX(B{r}+E{r}-D{r},0)",
                   total="=SUM(F{f}:F{l})", number_format=MONEY, width=20),
        ),
        howto=("Type each debt's current balance, its yearly interest rate and what you "
               "pay each month.",
               "Interest this month and Balance after payment work themselves out.",
               "Each month, copy Balance after payment into Balance.",
               "The Example sheet shows the same layout with sample numbers.",
               "Rename any debt by typing over it."),
    ),
}


def kind_for(keyword: str) -> str | None:
    """The kind a keyword asks for: it names a fill-in format and shares a word with
    exactly one kind's vocabulary (the first by most shared words). None: we cannot make it."""
    words = set(re.findall(r"[a-z]+", keyword.lower()))
    if not words & FORMAT_WORDS:
        return None
    scored = sorted(((len(words & k.vocabulary), name) for name, k in KINDS.items()),
                    reverse=True)
    best, name = scored[0]
    return name if best > 0 and (len(scored) == 1 or scored[1][0] < best) else None


# ---- the layout -------------------------------------------------------------------------
@dataclass(frozen=True)
class Layout:
    header_row: int
    first: int
    last: int
    total: int
    columns: tuple


def layout_of(kind: Kind, n_labels: int) -> Layout:
    header = 2 + len(kind.top) + 1
    first = header + 1
    last = first + n_labels - 1
    return Layout(header, first, last, last + 1, kind.columns)


def formulas_of(kind: Kind, n_labels: int) -> dict[str, str]:
    """Every formula cell -> its formula, as written in BOTH computed sheets."""
    lay = layout_of(kind, n_labels)
    out: dict[str, str] = {}
    fmt = {"f": lay.first, "l": lay.last, "t": lay.total, "n": n_labels}
    for top in kind.top:
        if top.role == "formula":
            out[top.cell] = top.formula.format(**fmt)
    for c, col in enumerate(kind.columns, 1):
        letter = col_letters(c)
        if col.role == "formula":
            for r in range(lay.first, lay.last + 1):
                out[f"{letter}{r}"] = col.formula.format(r=r)
        if col.total:
            out[f"{letter}{lay.total}"] = col.total.format(**fmt)
    return out


def build_workbook(kind_name: str, headline: str, labels: list[str]) -> bytes:
    """The .xlsx bytes. Same cells for the same inputs; the bytes themselves differ by the
    zip's entry times, which is why every staged file is pinned by its SHA-256."""
    if openpyxl_problem():
        raise OpenpyxlMissing(openpyxl_problem())
    import datetime as _dt

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.properties import PageSetupProperties

    kind = KINDS[kind_name]
    lay = layout_of(kind, len(labels))
    cells = formulas_of(kind, len(labels))
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="16202A")
    calc_fill = PatternFill("solid", fgColor="E8F4EF")
    total_fill = PatternFill("solid", fgColor="D3E9E0")
    for index, name in enumerate(SHEETS[:2]):
        ws = wb.active if index == 0 else wb.create_sheet(name)
        ws.title = name
        example = name == "Example"
        ws.cell(TITLE_ROW, 1, headline + (" - example" if example else "")).font = \
            Font(bold=True, size=16)
        for top in kind.top:
            row = int(re.sub(r"\D", "", top.cell))
            ws.cell(row, 1, top.label).font = Font(bold=True)
            if top.role == "formula":
                ws[top.cell] = cells[top.cell]
                ws[top.cell].fill = calc_fill
            else:
                value = top.sample if example else top.value
                if value is not None:
                    ws[top.cell] = value
            if top.number_format:
                ws[top.cell].number_format = top.number_format
        for c, col in enumerate(kind.columns, 1):
            letter = col_letters(c)
            h = ws.cell(lay.header_row, c, col.header)
            h.font = Font(bold=True, color="FFFFFF")
            h.fill = head_fill
            h.alignment = Alignment(horizontal="center", wrap_text=True)
            ws.column_dimensions[letter].width = col.width
            for i, r in enumerate(range(lay.first, lay.last + 1)):
                cell = ws.cell(r, c)
                if col.role == "label":
                    cell.value = labels[i]
                elif col.role == "formula":
                    cell.value = cells[f"{letter}{r}"]
                    cell.fill = calc_fill
                elif example and col.sample is not None:
                    cell.value = col.sample(i)
                if col.number_format:
                    cell.number_format = col.number_format
            t = ws.cell(lay.total, c)
            if c == 1:
                t.value = "Total"
            elif col.total:
                t.value = cells[f"{letter}{lay.total}"]
            t.font = Font(bold=True)
            t.fill = total_fill
            if col.number_format:
                t.number_format = col.number_format
        ws.freeze_panes = ws.cell(lay.first, 2)
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation = "landscape" if kind.landscape else "portrait"
    howto = wb.create_sheet(SHEETS[2])
    howto.cell(1, 1, f"How to use your {kind.noun}").font = Font(bold=True, size=14)
    for i, line in enumerate(kind.howto, 3):
        howto.cell(i, 1, line)
    howto.column_dimensions["A"].width = 100
    stamp = _dt.datetime(2026, 1, 1)
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.properties.creator = "shop owner"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- reading it back ----------------------------------------------------------------------
@dataclass
class Grid:
    """One sheet as read back from the file: display text per cell, and where formulas are."""

    sheet: str
    rows: list                     # list of lists of str
    formulas: dict                 # "B5" -> "=..."
    values: dict                   # "B5" -> computed value
    header_row: int = 0
    total_row: int = 0
    width: int = 0


def _show(value, number_format: str) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, (int, float)):
        if "%" in number_format:
            return f"{value * 100:.0f}%"
        if number_format == MONEY:
            return f"{value:,.2f}"
        return f"{value:g}" if isinstance(value, float) else str(value)
    return str(value)


def read_sheet(xlsx: bytes, sheet: str) -> Grid:
    """The sheet's cells from the .xlsx bytes, formulas computed with ``formulas``."""
    if openpyxl_problem():
        raise OpenpyxlMissing(openpyxl_problem())
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx), data_only=False)
    ws = wb[sheet]
    raw: dict[str, object] = {}
    fmts: dict[str, str] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                raw[cell.coordinate] = cell.value
            fmts[cell.coordinate] = cell.number_format or ""
    ev = Evaluator(lambda ref: raw.get(ref))
    values = {ref: ev.value(ref) for ref in raw}
    rows = []
    for r in range(1, ws.max_row + 1):
        rows.append([_show(values.get(f"{col_letters(c)}{r}"),
                           fmts.get(f"{col_letters(c)}{r}", ""))
                     for c in range(1, ws.max_column + 1)])
    formulas = {k: v for k, v in raw.items() if isinstance(v, str) and v.startswith("=")}
    header = next((r for r in range(1, ws.max_row + 1)
                   if ws.cell(r, 1).value in {k.columns[0].header for k in KINDS.values()}), 0)
    total = next((r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Total"), 0)
    return Grid(sheet, rows, formulas, values, header, total, ws.max_column)


def sheet_names(xlsx: bytes) -> list[str]:
    from openpyxl import load_workbook
    return list(load_workbook(io.BytesIO(xlsx), read_only=True).sheetnames)


def howto_lines(xlsx: bytes) -> list[str]:
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(xlsx))[SHEETS[2]]
    return [str(c.value) for c in ws["A"] if c.value]
